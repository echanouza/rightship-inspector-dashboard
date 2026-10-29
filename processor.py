from pathlib import Path
import re
import pandas as pd

SUPPORTED = {".xlsx", ".xls", ".csv", ".pdf", ".docx", ".doc"}

def extract_text(path: Path):
    ext = path.suffix.lower()
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception as e:
            raise RuntimeError(f"PDF text extraction failed: {e}")
    if ext == ".docx":
        try:
            from docx import Document
            doc = Document(str(path))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(parts)
        except Exception as e:
            raise RuntimeError(f"Word extraction failed: {e}")
    if ext == ".doc":
        raise RuntimeError("Legacy .doc files are not supported by this starter. Save as .docx or PDF.")
    return ""

def norm_col(value):
    return re.sub(r"[^a-z0-9]+", "", str(value).strip().lower())

ALIASES = {
    "inspector": ["inspector", "inspectorname", "rightshipinspector", "surveyorauditor", "auditor"],
    "location": ["location", "inspectionlocation", "port", "portofinspection", "inspectionport"],
    "vessel": ["vessel", "vesselname", "ship", "shipname"],
    "inspection_date": ["inspectiondate", "dateofinspection", "date", "inspectioncompleted"],
    "finding_count": ["findingcount", "numberoffindings", "totalfindings", "observations", "numberofobservations", "totalobservations", "observationscount"],
    "report_id": ["reportid", "inspectionid", "reference", "reportnumber"]
}

def map_columns(df):
    lookup = {norm_col(c): c for c in df.columns}
    mapping = {}
    for target, aliases in ALIASES.items():
        for alias in aliases:
            if alias in lookup:
                mapping[target] = lookup[alias]
                break
    return mapping

def count_findings_from_text(text):
    patterns = [
        r"(?:total\s+)?(?:number\s+of\s+)?(?:findings|observations)\s*[:\-]\s*(\d+)",
        r"(\d+)\s+(?:findings|observations)\b"
    ]
    for pattern in patterns:
        m = re.search(pattern, text, re.I)
        if m:
            return int(m.group(1))
    # Conservative fallback: do not guess a count from unstructured text.
    return None

def parse_file(path):
    ext = path.suffix.lower()
    records, findings = [], []
    if ext in {".xlsx", ".xls", ".csv"}:
        sheets = pd.read_excel(path, sheet_name=None) if ext in {".xlsx", ".xls"} else {"CSV": pd.read_csv(path)}
        for sheet_name, df in sheets.items():
            if df.empty:
                continue
            mapping = map_columns(df)
            if not any(k in mapping for k in ["inspector", "location", "vessel", "finding_count"]):
                continue
            for _, row in df.iterrows():
                rec = {k: row[v] for k, v in mapping.items()}
                rec.update({"source_file": path.name, "source_section": sheet_name})
                records.append(rec)
    elif ext in {".pdf", ".docx"}:
        text = extract_text(path)
        if not text.strip():
            raise RuntimeError("No selectable text found; this may be a scanned document requiring OCR.")
        inspector = re.search(r"(?:rightship\s+)?inspector(?:\s+name)?\s*[:\-]\s*([^\n\r]+)", text, re.I)
        location = re.search(r"(?:inspection\s+location|port\s+of\s+inspection|location)\s*[:\-]\s*([^\n\r]+)", text, re.I)
        vessel = re.search(r"(?:vessel\s+name|ship\s+name|vessel)\s*[:\-]\s*([^\n\r]+)", text, re.I)
        date = re.search(r"(?:inspection\s+date|date\s+of\s+inspection)\s*[:\-]\s*([^\n\r]+)", text, re.I)
        rec = {
            "inspector": inspector.group(1).strip() if inspector else None,
            "location": location.group(1).strip() if location else None,
            "vessel": vessel.group(1).strip() if vessel else None,
            "inspection_date": date.group(1).strip() if date else None,
            "finding_count": count_findings_from_text(text),
            "source_file": path.name,
            "source_section": "Document"
        }
        if any(rec.get(k) for k in ["inspector", "location", "vessel", "finding_count"]):
            records.append(rec)
    return records, findings

def load_reports(folder: Path):
    warnings, file_rows, records, findings = [], [], [], []
    if not folder.exists():
        return pd.DataFrame(), pd.DataFrame(), pd.DataFrame(), [f"Folder does not exist: {folder}"]
    for path in sorted(folder.rglob("*")):
        if not path.is_file() or path.suffix.lower() not in SUPPORTED or path.name.startswith("~$"):
            continue
        file_rows.append({"file": str(path), "type": path.suffix.lower(), "size_kb": round(path.stat().st_size / 1024, 1)})
        try:
            recs, fins = parse_file(path)
            records.extend(recs)
            findings.extend(fins)
            if not recs:
                warnings.append(f"{path.name}: no inspection record could be confidently extracted; check the report layout.")
        except Exception as e:
            warnings.append(f"{path.name}: {e}")
    inspections = pd.DataFrame(records)
    if not inspections.empty:
        for col in ["inspector", "location", "vessel", "inspection_date", "finding_count", "report_id", "source_file", "source_section"]:
            if col not in inspections:
                inspections[col] = None
        inspections["finding_count"] = pd.to_numeric(inspections["finding_count"], errors="coerce")
        inspections["inspection_date"] = pd.to_datetime(inspections["inspection_date"], errors="coerce")
        for col in ["inspector", "location", "vessel"]:
            inspections[col] = inspections[col].astype("string").str.strip()
    return inspections, pd.DataFrame(findings), pd.DataFrame(file_rows), warnings
